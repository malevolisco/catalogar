# -*- coding: utf-8 -*-
"""
listas.py - Lector de listas de trabajo (los Search Results de MediaCentral).

Una lista de trabajo es el CSV o el Excel que se exporta de la busqueda de MediaCentral y que dice
que envios de agencia hay pendientes. No es el Excel de fichas de excel.py, que es otra cosa: ese
trae las fichas ya escritas para completarlas, este solo trae los numeros que hay que catalogar.

Formato tipico:

    Name,Catalogador,Comments,Notas Internas,Created,
    "4359 USA-MAMDANI LULA-PRESSER: MAMDANI PRAISES LULA'S ...",,"TRATAR AGENCIAS",,"21.09.2026 22:15:20 +02:00",
    "4685905 US TX ICE AGENT SHOT PRESSER 20260921I: LULAC ...",,"TRATAR AGENCIAS",,"21.09.2026 22:54:37 +02:00",

Reglas, en este orden:
  1. El numero de envio es el que abre la columna Name. Lo que va detras es el slug y el titular de
     la agencia: se guarda solo como etiqueta para reconocer la fila en el panel, nunca como fuente
     para redactar, que la ficha se lee siempre de la web de la agencia.
  2. Fila con algo en la columna Catalogador: si ese usuario de MediaCentral esta en "catalogadores"
     de config.json ({"I23785": "yo", "I24082": "javier"}), la fila se hace y sale con "para": a
     quien hay que mandarle la ficha ("yo" no se manda: se ve en la pagina). Si no se sabe quien es,
     fuera, y se dice el usuario para poder añadirlo.
  3. En la columna de fecha (Created) nunca se buscan numeros: su 2026 es el error clasico de
     rastrear el fichero entero buscando cuatro cifras. Lo que si se lee de ella es cuando entro el
     envio en MediaCentral ("creado": dd/mm/aaaa HH:MM), porque los numeros de Reuters se repiten
     cada pocos dias y ese momento dice cual de ellos es: el ultimo que la agencia mando antes.
  4. Numeros repetidos, una sola vez, en el orden en que aparecen.

Uso:
    from listas import leer_lista, leer_texto, ListaError
    envios, descartes, aviso = leer_lista("Search Results.csv", datos_en_bytes)   # fichero de MediaCentral
    trabajos, ignorados = leer_texto("hazme el 7612, 7613 y AP4681323")          # texto pegado de cualquier forma
"""
import csv
import io
import re

# Los mismos tres formatos que entiende el resto de la herramienta: Reuters, AP y EBU.
# Al principio de la fila se admiten tambien tres cifras (624), que es el mismo envio que 0624 y
# que algun export deja sin el cero; sueltas dentro del texto no, que ahi cualquier cifra vale.
NUMERO = r"(?:AP)?(\d{4}_\d{6,9}|\d{7}|\d{4})"
NUMERO_INICIAL_RE = re.compile(r"^\s*(?:AP)?(\d{4}_\d{6,9}|\d{7}|\d{3,4})(?![\d/_-])", re.I)
NUMERO_SUELTO_RE = re.compile(r"(?<![\d/_-])" + NUMERO + r"(?![\d/_-])", re.I)

# Cabeceras que se reconocen, en minusculas y sin acentos (se comparan ya normalizadas).
CAB_TITULO = ("name", "nombre", "titulo", "title", "asunto", "envio", "descripcion")
CAB_CATALOGADOR = ("catalogador", "cataloged by", "cataloger", "asignado", "usuario")
CAB_IGNORAR = ("created", "creado", "fecha", "date", "modified", "modificado", "duracion", "duration")
CAB_CREADO = ("created", "creado", "fecha de creacion", "creation date")

# "26.09.2026 01:40:28 +02:00", "26/09/2026 1:40", "2026-09-26 01:40:28": fecha y hora de entrada en MediaCentral
CREADO_RE = re.compile(r"(?<!\d)(\d{1,2})[./-](\d{1,2})[./-](\d{4})[ T]+(\d{1,2}):(\d{2})")
CREADO_ISO_RE = re.compile(r"(?<!\d)(\d{4})-(\d{2})-(\d{2})[ T]+(\d{1,2}):(\d{2})")


def leer_creado(texto):
    """La fecha y hora de la columna Created como "dd/mm/aaaa HH:MM", o "" si no la hay."""
    m = CREADO_RE.search(texto or "")
    if m:
        d, mes, a, h, mi = m.groups()
    else:
        m = CREADO_ISO_RE.search(texto or "")
        if not m:
            return ""
        a, mes, d, h, mi = m.groups()
    try:
        return f"{int(d):02d}/{int(mes):02d}/{a} {int(h):02d}:{mi}"
    except ValueError:
        return ""

MAX_FILAS = 2000            # un export de MediaCentral no pasa de unos cientos
LARGO_ETIQUETA = 120        # lo que se guarda del titular para reconocer la fila
CFG = {"catalogadores": {}}   # usuario de MediaCentral (sin mayusculas) -> persona de la agenda, "yo" o direccion


def configurar(cfg):
    """Toma de config.json la tabla "catalogadores": {"I23785": "yo", "I24082": "javier"}."""
    tabla = cfg.get("catalogadores") or {}
    CFG["catalogadores"] = {_sin_acentos(k): str(v).strip() for k, v in tabla.items() if str(v).strip()}


def persona_de(catalogador):
    """A quien corresponde un usuario de MediaCentral, o "" si no esta en la tabla."""
    return CFG["catalogadores"].get(_sin_acentos(catalogador or ""), "")


class ListaError(Exception):
    pass


def _sin_acentos(t):
    pares = (("á", "a"), ("é", "e"), ("í", "i"), ("ó", "o"), ("ú", "u"), ("ü", "u"), ("ñ", "n"))
    t = (t or "").strip().lower()
    for a, b in pares:
        t = t.replace(a, b)
    return t


def _decodificar(datos):
    """Texto de un fichero de texto venga como venga de Windows, de Gmail o de MediaCentral."""
    if isinstance(datos, str):
        return datos
    for cod in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return datos.decode(cod)
        except UnicodeDecodeError:
            continue
    return datos.decode("utf-8", "replace")


def _filas_csv(datos):
    """Filas de un CSV, con el separador que traiga (coma o punto y coma)."""
    texto = _decodificar(datos)
    if not texto.strip():
        raise ListaError("El fichero esta vacio")
    muestra = texto[:4000]
    try:
        dialecto = csv.Sniffer().sniff(muestra, delimiters=",;\t")
    except csv.Error:
        dialecto = csv.excel                      # coma, que es lo que exporta MediaCentral
    return [f for f in csv.reader(io.StringIO(texto), dialecto) if any((c or "").strip() for c in f)]


def _filas_xlsx(datos):
    """Filas de la primera hoja de un Excel, como texto."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ListaError("Falta openpyxl para leer ficheros .xlsx (pip install openpyxl)")
    try:
        libro = load_workbook(io.BytesIO(datos), read_only=True, data_only=True)
    except Exception as e:
        raise ListaError(f"No se pudo abrir el Excel ({type(e).__name__})")
    filas = []
    for fila in libro.worksheets[0].iter_rows(values_only=True):
        valores = ["" if c is None else str(c).strip() for c in fila]
        if any(valores):
            filas.append(valores)
        if len(filas) > MAX_FILAS + 1:
            break
    libro.close()
    if not filas:
        raise ListaError("La primera hoja del Excel esta vacia")
    return filas


def _columnas(cabecera):
    """Devuelve (indice del titulo, del catalogador, de Created, indices a ignorar, hay cabecera) leyendo
    la cabecera. Si no hay cabecera reconocible, el titulo es la primera columna y no se ignora nada."""
    normal = [_sin_acentos(c) for c in cabecera]
    i_titulo = next((i for i, c in enumerate(normal) if c in CAB_TITULO), None)
    i_catalogador = next((i for i, c in enumerate(normal) if c in CAB_CATALOGADOR), None)
    i_creado = next((i for i, c in enumerate(normal) if c in CAB_CREADO), None)
    ignorar = {i for i, c in enumerate(normal) if c in CAB_IGNORAR}
    hay_cabecera = i_titulo is not None or i_catalogador is not None
    if i_titulo is None:
        i_titulo = 0
    return i_titulo, i_catalogador, i_creado, ignorar, hay_cabecera


def _normalizar(numero):
    """0624 y 624 son el mismo envio de Reuters; los de AP y EBU se quedan como estan."""
    return numero.zfill(4) if len(numero) <= 4 else numero


def leer_lista(nombre, datos):
    """Lee una lista de trabajo y devuelve (envios, descartes, aviso).

    envios:    [{"numero": "4359", "etiqueta": "USA-MAMDANI LULA-PRESSER: ...", "fila": 2,
                 "catalogador": "I24082", "para": "javier", "creado": "21/09/2026 22:15"}, ...]
               (para: a quien mandarle la ficha, o ""; creado: cuando entro en MediaCentral, o "")
    descartes: [{"fila": 3, "texto": "...", "motivo": "ya la tiene I99999"}, ...]
    aviso:     "" o una frase corta que conviene ensenar (por ejemplo, que el numero no abria la fila)
    """
    nombre = (nombre or "").lower()
    if nombre.endswith(".xlsx"):
        filas = _filas_xlsx(datos)
    elif nombre.endswith((".csv", ".txt", ".tsv")):
        filas = _filas_csv(datos)
    else:
        raise ListaError("Solo se admiten ficheros .csv, .txt o .xlsx")
    if not filas:
        raise ListaError("El fichero no tiene ninguna fila con datos")

    i_titulo, i_catalogador, i_creado, ignorar, hay_cabecera = _columnas(filas[0])
    cuerpo = filas[1:] if hay_cabecera else filas
    if len(cuerpo) > MAX_FILAS:
        raise ListaError(f"La lista tiene {len(cuerpo)} filas y el tope son {MAX_FILAS}")

    # Primera pasada: el numero abre la columna del titulo, que es como los exporta MediaCentral.
    envios, descartes, vistos = [], [], set()
    sin_numero_al_principio = []
    catalogadores = {}                     # quien tiene cada fila descartada, para el mensaje final
    primera = 2 if hay_cabecera else 1

    for n, fila in enumerate(cuerpo, start=primera):
        def celda(i):
            return (fila[i].strip() if i is not None and i < len(fila) and fila[i] else "")

        texto = celda(i_titulo)
        if not texto:
            continue
        quien = celda(i_catalogador)
        para = persona_de(quien)
        if quien and not para:
            descartes.append({"fila": n, "texto": texto[:LARGO_ETIQUETA], "motivo": f"ya la tiene {quien}"})
            catalogadores[quien] = catalogadores.get(quien, 0) + 1
            continue
        m = NUMERO_INICIAL_RE.match(texto)
        if not m:
            sin_numero_al_principio.append((n, fila, texto))
            continue
        numero = _normalizar(m.group(1))
        etiqueta = texto[m.end():].strip(" -–—:·.")[:LARGO_ETIQUETA]
        if numero in vistos:
            descartes.append({"fila": n, "texto": etiqueta, "motivo": f"{numero} repetido"})
            continue
        vistos.add(numero)
        envios.append({"numero": numero, "etiqueta": etiqueta, "fila": n, "catalogador": quien, "para": para,
                       "creado": leer_creado(celda(i_creado))})

    aviso = ""
    # Segunda pasada, solo si ninguna fila traia el numero delante: se busca dentro del titulo,
    # nunca en la columna de fecha, y se avisa de que se ha hecho asi.
    if not envios and sin_numero_al_principio:
        for n, fila, texto in sin_numero_al_principio:
            candidatos = [c for i, c in enumerate(fila) if i not in ignorar and i != i_catalogador]
            m = NUMERO_SUELTO_RE.search(" ".join(candidatos))
            if not m:
                descartes.append({"fila": n, "texto": texto[:LARGO_ETIQUETA], "motivo": "sin numero de envio"})
                continue
            numero = _normalizar(m.group(1))
            if numero in vistos:
                descartes.append({"fila": n, "texto": texto[:LARGO_ETIQUETA], "motivo": f"{numero} repetido"})
                continue
            vistos.add(numero)
            envios.append({"numero": numero, "etiqueta": texto[:LARGO_ETIQUETA], "fila": n, "catalogador": "", "para": "",
                           "creado": leer_creado(fila[i_creado] if i_creado is not None and i_creado < len(fila) else "")})
        if envios:
            aviso = ("En esta lista el numero no abria la fila: se ha cogido el primero que aparece en cada una. "
                     "Conviene comprobarlos antes de lanzar el lote.")
    else:
        for n, _fila, texto in sin_numero_al_principio:
            descartes.append({"fila": n, "texto": texto[:LARGO_ETIQUETA], "motivo": "sin numero de envio"})

    if not envios:
        if catalogadores and not sin_numero_al_principio:
            quienes = ", ".join(f"{k} ({v})" for k, v in catalogadores.items())
            raise ListaError(f"Todas las filas de la lista tienen catalogador y no se quien es: {quienes}. "
                             f"Añadelo a \"catalogadores\" en config.json, por ejemplo "
                             f"{{\"{next(iter(catalogadores))}\": \"yo\"}} o {{\"{next(iter(catalogadores))}\": \"andrea\"}}, "
                             "y sus filas se haran y le llegaran por correo")
        raise ListaError("Ninguna fila de la lista lleva numero de envio de agencia "
                         "(Reuters de 4 cifras, AP de 7 o EBU tipo 2026_10420363)")
    descartes.sort(key=lambda d: d["fila"])
    return envios, descartes, aviso


def resumen(envios, descartes):
    """Una linea para el registro o para el correo."""
    texto = f"{len(envios)} envio(s)"
    if descartes:
        motivos = {}
        for d in descartes:
            clave = "repetidos" if "repetido" in d["motivo"] else \
                    "ya catalogadas" if "ya la tiene" in d["motivo"] else "sin numero"
            motivos[clave] = motivos.get(clave, 0) + 1
        texto += ", " + ", ".join(f"{v} {k}" for k, v in motivos.items())
    return texto


# ====================================================================== numeros sueltos en un texto
# Lo que se pega en la caja de numeros, lo que llega en el cuerpo de un correo o los argumentos de
# la consola: vale cualquier cosa ("3213 2314", "3213,2314", "3242 letras", "AP4681323", una linea
# con texto alrededor). Se queda solo con los numeros de envio y sabe de que agencia es cada uno.

# Por orden: EBU, fechas (dd/mm/aaaa con / . o -, y aaaa-mm-dd), horas, y cualquier tira de cifras.
_PIEZA_RE = re.compile(
    r"(?P<ebu>(?<![\d_])\d{4}_\d{6,9}(?![\d_]))"
    r"|(?P<fecha>(?<!\d)\d{1,2}[/.-]\d{1,2}[/.-]\d{4}(?!\d))"
    r"|(?P<iso>(?<!\d)\d{4}-\d{1,2}-\d{1,2}(?!\d))"
    r"|(?P<hora>(?<!\d)\d{1,2}:\d{2}(?::\d{2})?(?!\d))"
    r"|(?P<cifras>\d+)"
)
# Lo que puede haber entre un numero y la fecha que se le aplica: "0624=06/09/2026", "0624 del 06/09/2026".
_ENLACE_FECHA_RE = re.compile(r"^\s*(?:[=@:,]|del?|el|en)?\s*$", re.I)


def agencia_de(numero):
    """Reuters (4 cifras), AP (7) o EBU (2026_10420363): lo mismo que decide extractor.fetch."""
    if "_" in numero:
        return "EBU"
    return "AP" if len(numero) == 7 else "Reuters"


def leer_texto(texto):
    """Numeros de envio de un texto escrito de cualquier manera.

    Devuelve (trabajos, ignorados):
      trabajos  = [(numero, fecha o None), ...] sin repetir y en el orden en que aparecen
      ignorados = cifras que parecian un numero pero no lo son, con el motivo, para decirlo en pantalla

    Una fecha pegada detras de un numero ("0624 06/09/2026", "0624=06/09/2026") es la fecha de ese
    envio. Las demas fechas y las horas se saltan: su 2026 no es un envio de Reuters.
    """
    trabajos, ignorados = [], []
    ultimo_fin = None          # donde acaba el ultimo numero sin fecha, por si le sigue una
    for m in _PIEZA_RE.finditer(texto or ""):
        tipo, valor = m.lastgroup, m.group()
        if tipo in ("fecha", "iso"):
            if ultimo_fin is not None and _ENLACE_FECHA_RE.match(texto[ultimo_fin:m.start()]):
                if tipo == "iso":
                    a, mes, d = valor.split("-")
                else:
                    d, mes, a = re.split(r"[/.-]", valor)
                trabajos[-1] = (trabajos[-1][0], f"{int(d):02d}/{int(mes):02d}/{a}")
            ultimo_fin = None
            continue
        if tipo == "hora":
            ultimo_fin = None
            continue
        if tipo == "cifras":
            n = len(valor)
            if n not in (4, 7):
                if n >= 3:
                    pista = " (si es de Reuters, con el cero delante: 0" + valor + ")" if n == 3 else ""
                    ignorados.append(f"{valor}: {n} cifras{pista}")
                ultimo_fin = None
                continue
        if any(numero == valor for numero, _ in trabajos):      # repetido: una sola vez
            ultimo_fin = None
            continue
        trabajos.append((valor, None))
        ultimo_fin = m.end()
    return trabajos, ignorados


def resumen_agencias(numeros):
    """ "3 de Reuters, 1 de AP" para decir de un vistazo a donde va cada uno."""
    cuenta = {}
    for n in numeros:
        a = agencia_de(n)
        cuenta[a] = cuenta.get(a, 0) + 1
    return ", ".join(f"{v} de {k}" for k, v in cuenta.items())
