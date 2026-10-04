# -*- coding: utf-8 -*-
"""
excel.py - Entrada y salida por Excel para catalogar.py (--excel).

Entrada: un .xlsx con cabecera en alguna de las diez primeras filas y, como minimo,
las columnas NAME (nombre original del envio) y CONT (script original). Opcionales:
ENVIO o NUMERO (numero de envio, solo para fotogramas), FECHA y SLUG.
Los nombres de columna se reconocen sin distinguir mayusculas, tildes ni puntos.

Salida: copia del Excel con cuatro columnas anadidas al final de la cabecera
(NAME_ARCHIVO, COMMENT, RESTRICCIONES, AVISOS), guardada como <nombre>_catalogado.xlsx.
El original no se toca. Se guarda tras cada fila, asi que un corte no pierde nada,
y con --reanudar se saltan las filas que ya tienen NAME_ARCHIVO.
"""
import re
import unicodedata
from datetime import datetime, date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

COLUMNAS_ENTRADA = {
    "headline": ("NAME", "NOMBRE", "TITULO", "TITLE", "HEADLINE", "TITULO PARALELO"),
    "texto":    ("CONT", "SCRIPT", "TEXTO", "CONTENIDO", "SHOTLIST", "CONT ORIGINAL"),
    "numero":   ("ENVIO", "NUMERO", "NUM", "N", "NO", "EDIT NO", "STORY NO", "ID", "NUMERO DE ENVIO"),
    "fecha":    ("FECHA", "DATE"),
    "slug":     ("SLUG",),
}
COLUMNAS_SALIDA = ("NAME_ARCHIVO", "COMMENT", "RESTRICCIONES", "AVISOS")
ANCHOS_SALIDA = {"NAME_ARCHIVO": 45, "COMMENT": 80, "RESTRICCIONES": 30, "AVISOS": 50}
FILAS_CABECERA = 10       # se busca la cabecera en las diez primeras filas


class ExcelError(Exception):
    pass


def _norm(valor):
    """Nombre de columna comparable: sin tildes, mayusculas, sin puntos ni dobles espacios."""
    s = "" if valor is None else str(valor)
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[.:_\-]+", " ", s.upper())
    return re.sub(r"\s+", " ", s).strip()


def _numero(valor):
    """Numero de envio como texto: 624 / 624.0 / '0624' / 'AP4681323' -> '0624' / '4681323'. '' si no vale."""
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    s = str(valor).strip().upper()
    if s.startswith("AP"):
        s = s[2:].strip()
    if not s.isdigit():
        return ""
    if len(s) <= 4:
        return s.zfill(4)         # Excel quita los ceros por la izquierda: 624 -> 0624
    return s


def _fecha(valor):
    """Fecha como DD/MM/AAAA si se reconoce; si no, tal cual."""
    if valor is None:
        return ""
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%d/%m/%Y")
    s = str(valor).strip()
    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})(?:[ T].*)?", s)
    if m:
        return f"{m.group(3)}/{m.group(2)}/{m.group(1)}"
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        return f"{int(m.group(1)):02d}/{int(m.group(2)):02d}/{m.group(3)}"
    return s


def _texto(valor):
    if valor is None:
        return ""
    return str(valor).replace("\r\n", "\n").strip()


class LibroExcel:
    """Abre el Excel de entrada (o el de salida si se reanuda), localiza las columnas
    y escribe las fichas fila a fila."""

    def __init__(self, origen, destino=None, reanudar=False):
        self.origen = Path(origen)
        if not self.origen.exists():
            raise ExcelError(f"No existe el fichero {self.origen}")
        if self.origen.suffix.lower() != ".xlsx":
            raise ExcelError("Solo se admite .xlsx (guarda el fichero como Libro de Excel)")
        self.destino = Path(destino) if destino else self.origen.with_name(self.origen.stem + "_catalogado.xlsx")
        if self.destino.resolve() == self.origen.resolve():
            raise ExcelError("El fichero de salida no puede ser el mismo que el de entrada")
        self.reanudado = bool(reanudar and self.destino.exists())
        fuente = self.destino if self.reanudado else self.origen
        try:
            self.wb = load_workbook(fuente)
        except Exception as e:
            raise ExcelError(f"No se puede abrir {fuente}: {e}")
        self.ws = self.wb.active
        self.fila_cabecera, self.col = self._localizar_cabecera()
        self.col_salida = self._preparar_salida()

    # ------------------------------------------------------------------ cabecera
    def _localizar_cabecera(self):
        ws = self.ws
        for r in range(1, min(FILAS_CABECERA, ws.max_row) + 1):
            encontradas = {}
            for c in range(1, ws.max_column + 1):
                nombre = _norm(ws.cell(row=r, column=c).value)
                if not nombre:
                    continue
                for clave, alias in COLUMNAS_ENTRADA.items():
                    if nombre in alias and clave not in encontradas:
                        encontradas[clave] = c
                for salida in COLUMNAS_SALIDA:
                    if nombre == _norm(salida):
                        encontradas["_" + salida] = c
            if "headline" in encontradas and "texto" in encontradas:
                return r, encontradas
        raise ExcelError(
            "No encuentro la cabecera: hace falta una fila con las columnas NAME y CONT "
            "(y, si se quiere, ENVIO, FECHA y SLUG). Renombra la primera fila y repite.")

    def _preparar_salida(self):
        """Devuelve {nombre_columna_salida: indice}. Las crea al final de la cabecera si no existen."""
        ws = self.ws
        col = {}
        siguiente = ws.max_column + 1
        for nombre in COLUMNAS_SALIDA:
            c = self.col.get("_" + nombre)
            if c is None:
                c = siguiente
                siguiente += 1
                celda = ws.cell(row=self.fila_cabecera, column=c, value=nombre)
                celda.font = Font(bold=True)
                ws.column_dimensions[get_column_letter(c)].width = ANCHOS_SALIDA[nombre]
            col[nombre] = c
        return col

    # ------------------------------------------------------------------ lectura
    def filas(self):
        """Lista de dicts: fila, ficha, hecha, motivo (por que se salta, o '')."""
        ws = self.ws
        resultado = []
        for r in range(self.fila_cabecera + 1, ws.max_row + 1):
            def v(clave):
                c = self.col.get(clave)
                return ws.cell(row=r, column=c).value if c else None
            headline = _texto(v("headline"))
            texto = _texto(v("texto"))
            numero = _numero(v("numero"))
            if not headline and not texto and not numero:
                continue                                   # fila vacia
            hecha = bool(_texto(ws.cell(row=r, column=self.col_salida["NAME_ARCHIVO"]).value))
            motivo = ""
            if not texto:
                motivo = "sin CONT (script)"
            elif not headline:
                motivo = "sin NAME (nombre original)"
            ficha = {
                "agencia": "AP" if len(numero) >= 6 else ("REUTERS" if numero else ""),
                "numero": numero,
                "fecha": _fecha(v("fecha")),
                "rev": "",
                "slug": _texto(v("slug")),
                "headline": headline,
                "texto": texto,
                "url": "",
                "avisos": [],
            }
            resultado.append({"fila": r, "ficha": ficha, "hecha": hecha, "motivo": motivo})
        return resultado

    # ------------------------------------------------------------------ escritura
    def _poner(self, fila, nombre, valor):
        celda = self.ws.cell(row=fila, column=self.col_salida[nombre], value=valor)
        celda.alignment = Alignment(wrap_text=True, vertical="top")

    def escribir(self, fila, campos, avisos):
        self._poner(fila, "NAME_ARCHIVO", campos.get("NAME", ""))
        self._poner(fila, "COMMENT", campos.get("COMMENT", ""))
        self._poner(fila, "RESTRICCIONES", campos.get("RESTRICCIONES", "SIN AVISO"))
        self._poner(fila, "AVISOS", " | ".join(avisos) if avisos else "")

    def escribir_error(self, fila, mensaje):
        """Deja NAME_ARCHIVO vacio (asi --reanudar la vuelve a intentar) y el error en AVISOS."""
        self._poner(fila, "NAME_ARCHIVO", "")
        self._poner(fila, "COMMENT", "")
        self._poner(fila, "RESTRICCIONES", "")
        self._poner(fila, "AVISOS", "ERROR: " + mensaje)

    def guardar(self):
        """Guarda en destino. Si Excel tiene el fichero abierto (Windows), guarda en una copia con hora."""
        try:
            self.wb.save(self.destino)
            return self.destino
        except PermissionError:
            alternativo = self.destino.with_name(
                f"{self.destino.stem}_{datetime.now().strftime('%H%M%S')}.xlsx")
            self.wb.save(alternativo)
            print(f"AVISO: {self.destino.name} esta abierto en Excel; guardado en {alternativo.name}")
            self.destino = alternativo
            return alternativo
